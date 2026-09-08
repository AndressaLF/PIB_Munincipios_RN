"""Gera o PDF da atividade e o XLSX formatado com os dados solicitados."""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from src.pipeline import PASTA_DADOS, resumo_separatriz

AZUL = (31, 78, 121)
AZUL_CLARO = (217, 226, 243)
CINZA = (80, 80, 80)
BRANCO = (255, 255, 255)
VERMELHO = (153, 0, 0)

COLUNAS_PT = {
    "ranking_pib_per_capita_asc": "Ranking (menor PIB per capita)",
    "codigo_municipio": "Código IBGE",
    "municipio": "Município",
    "uf": "UF",
    "mesorregiao": "Mesorregião",
    "microrregiao": "Microrregião",
    "regiao_imediata": "Região imediata",
    "regiao_intermediaria": "Região intermediária",
    "ano_pib": "Ano do PIB",
    "pib_per_capita": "PIB per capita 2023 (R$)",
    "pib_mil_reais": "PIB 2023 (R$ mil)",
    "percentil_10_pib_per_capita": "Percentil 10 (R$)",
    "abaixo_percentil_10": "Abaixo do percentil 10",
    "entre_os_10_pct_menores": "Entre os 10% menores",
    "populacao_censo_2022": "População (Censo 2022)",
    "area_km2": "Área (km²)",
    "densidade_demografica": "Densidade (hab/km²)",
    "ano_vab": "Ano do VAB",
    "vab_agropecuaria_mil_reais": "VAB agropecuária 2021 (R$ mil)",
    "vab_industria_mil_reais": "VAB indústria 2021 (R$ mil)",
    "vab_servicos_mil_reais": "VAB serviços 2021 (R$ mil)",
    "vab_administracao_publica_mil_reais": "VAB administração pública 2021 (R$ mil)",
    "participacao_agropecuaria_pct": "Participação agropecuária no VAB (%)",
    "participacao_industria_pct": "Participação indústria no VAB (%)",
    "participacao_servicos_pct": "Participação serviços no VAB (%)",
    "participacao_adm_publica_pct": "Participação administração pública no VAB (%)",
    "taxa_alfabetizacao_15_mais_pct": "Taxa de alfabetização 15 anos ou mais (%)",
    "idhm": "IDHM",
    "taxa_escolarizacao_6_a_14_pct": "Taxa de escolarização 6 a 14 anos (%)",
    "salario_medio_mensal": "Salário médio mensal (salários mínimos)",
    "populacao_ocupada": "População ocupada (%)",
}

ORDEM_XLSX = [
    "ranking_pib_per_capita_asc",
    "codigo_municipio",
    "municipio",
    "uf",
    "mesorregiao",
    "microrregiao",
    "regiao_imediata",
    "pib_per_capita",
    "pib_mil_reais",
    "percentil_10_pib_per_capita",
    "abaixo_percentil_10",
    "populacao_censo_2022",
    "area_km2",
    "densidade_demografica",
    "participacao_agropecuaria_pct",
    "participacao_industria_pct",
    "participacao_servicos_pct",
    "participacao_adm_publica_pct",
    "vab_agropecuaria_mil_reais",
    "vab_industria_mil_reais",
    "vab_servicos_mil_reais",
    "vab_administracao_publica_mil_reais",
    "taxa_alfabetizacao_15_mais_pct",
    "idhm",
    "taxa_escolarizacao_6_a_14_pct",
    "salario_medio_mensal",
    "populacao_ocupada",
]


def br(valor, casas: int = 2) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return "-"
    texto = f"{float(valor):,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _fonte_windows() -> tuple[str, str]:
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    if regular.exists() and bold.exists():
        return str(regular), str(bold)
    raise FileNotFoundError("Fonte Arial não encontrada em C:\\Windows\\Fonts.")


class RelatorioPDF(FPDF):
    def __init__(self) -> None:
        super().__init__(orientation="P", unit="mm", format="A4")
        regular, bold = _fonte_windows()
        self.add_font("Texto", "", regular)
        self.add_font("Texto", "B", bold)
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(14, 18, 14)

    def header(self) -> None:
        if self.page_no() == 1:
            return
        self.set_font("Texto", "B", 9)
        self.set_text_color(*AZUL)
        self.cell(
            0,
            8,
            "Separatrizes - PIB per capita dos municípios do Rio Grande do Norte",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        self.set_draw_color(*AZUL)
        self.set_line_width(0.4)
        self.line(14, 16, 196, 16)
        self.ln(4)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Texto", "", 8)
        self.set_text_color(*CINZA)
        self.cell(0, 8, f"Página {self.page_no()}  |  Fonte: IBGE (Cidades@ e SIDRA)", align="C")

    def titulo(self, texto: str) -> None:
        self.set_font("Texto", "B", 13)
        self.set_text_color(*AZUL)
        self.multi_cell(0, 7, texto)
        self.ln(2)

    def corpo(self, texto: str) -> None:
        self.set_font("Texto", "", 10)
        self.set_text_color(40, 40, 40)
        self.multi_cell(0, 5.4, texto)
        self.ln(2)

    def destaque(self, texto: str) -> None:
        self.set_fill_color(*AZUL_CLARO)
        self.set_font("Texto", "B", 11)
        self.set_text_color(*AZUL)
        self.multi_cell(0, 8, f"  {texto}", fill=True)
        self.ln(3)

    def tabela(self, cabecalho: list[str], linhas: list[list[str]], larguras: list[float]) -> None:
        self.set_font("Texto", "B", 8)
        self.set_fill_color(*AZUL)
        self.set_text_color(*BRANCO)
        for i, col in enumerate(cabecalho):
            self.cell(larguras[i], 7, col, border=1, fill=True, align="C")
        self.ln()
        self.set_font("Texto", "", 7.5)
        self.set_text_color(30, 30, 30)
        for idx, linha in enumerate(linhas):
            if self.get_y() > 270:
                self.add_page()
                self.set_font("Texto", "B", 8)
                self.set_fill_color(*AZUL)
                self.set_text_color(*BRANCO)
                for i, col in enumerate(cabecalho):
                    self.cell(larguras[i], 7, col, border=1, fill=True, align="C")
                self.ln()
                self.set_font("Texto", "", 7.5)
                self.set_text_color(30, 30, 30)
            fill = idx % 2 == 1
            if fill:
                self.set_fill_color(242, 245, 250)
            for i, valor in enumerate(linha):
                align = "L" if i <= 2 else "R"
                if i == 0:
                    align = "C"
                self.cell(larguras[i], 5.6, str(valor), border=1, fill=fill, align=align)
            self.ln()
        self.ln(3)


def _grafico_histograma(df: pd.DataFrame, p10: float) -> BytesIO:
    valores = df["pib_per_capita"].dropna()
    fig, ax = plt.subplots(figsize=(8.4, 3.6), dpi=140)
    ax.hist(valores, bins=28, color="#4F81BD", edgecolor="white")
    ax.axvline(p10, color="#C00000", linestyle="--", linewidth=2, label=f"P10 = R$ {br(p10)}")
    ax.set_xlim(0, float(valores.quantile(0.95)))
    ax.set_xlabel("PIB per capita 2023 (R$)")
    ax.set_ylabel("Número de municípios")
    ax.set_title("Distribuição do PIB per capita municipal no RN")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _grafico_setores(resumo: dict) -> BytesIO:
    labels = ["Agropecuária", "Indústria", "Serviços", "Adm. pública"]
    chaves = [
        "Participação agropecuária no VAB (%)",
        "Participação indústria no VAB (%)",
        "Participação serviços no VAB (%)",
        "Participação administração pública no VAB (%)",
    ]
    mapa = {item["variavel"]: item for item in resumo["comparativo_grupo_vs_demais"]}
    grupo = [mapa[k]["media_grupo_p10"] for k in chaves]
    demais = [mapa[k]["media_demais_municipios"] for k in chaves]
    x = range(len(labels))
    fig, ax = plt.subplots(figsize=(8.4, 3.6), dpi=140)
    ax.bar([i - 0.18 for i in x], grupo, width=0.36, label="Grupo abaixo do P10", color="#C00000")
    ax.bar([i + 0.18 for i in x], demais, width=0.36, label="Demais municípios do RN", color="#4F81BD")
    ax.set_xticks(list(x))
    ax.set_xticklabels(labels)
    ax.set_ylabel("% do valor adicionado bruto (2021)")
    ax.set_title("Estrutura produtiva média: grupo P10 versus demais municípios")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    buf = BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def gerar_pdf(df: pd.DataFrame, resumo: dict, destino: Path) -> Path:
    pdf = RelatorioPDF()
    pdf.add_page()
    pdf.set_fill_color(*AZUL)
    pdf.rect(0, 0, 210, 42, "F")
    pdf.set_text_color(*BRANCO)
    pdf.set_font("Texto", "B", 18)
    pdf.set_xy(14, 10)
    pdf.multi_cell(182, 8, "Separatrizes: PIB per capita dos municípios do Rio Grande do Norte")
    pdf.set_font("Texto", "", 11)
    pdf.set_x(14)
    pdf.cell(0, 8, "Atividade de análise descritiva  |  IBGE Cidades@  |  Ano de referência: 2023")
    pdf.ln(16)

    pdf.titulo("Objetivo")
    pdf.corpo(
        "Este relatório responde à atividade de separatrizes aplicada ao PIB per capita "
        "de todos os municípios do Rio Grande do Norte. A separatriz utilizada é o "
        "percentil 10 (P10): o valor que deixa cerca de 10% das observações à esquerda "
        "da distribuição. Em seguida, os municípios com PIB per capita menor que o P10 "
        "são descritos com variáveis econômicas, demográficas e sociais que ajudam a "
        "explicar o resultado."
    )
    pdf.corpo(
        "Fonte dos dados: APIs oficiais do IBGE (Cidades@ / pesquisas e Agregados / SIDRA), "
        "as mesmas bases usadas no portal https://cidades.ibge.gov.br/brasil/rn/panorama. "
        "A listagem completa, com todas as variáveis, está no arquivo XLSX que acompanha este PDF."
    )

    pdf.titulo("1. PIB per capita de todos os municípios do RN")
    pdf.corpo(
        f"Foram obtidos os {resumo['n_municipios']} municípios do Rio Grande do Norte. "
        "O indicador é o PIB per capita da série revisada de 2023 (pesquisa 38, indicador 47001). "
        f"O menor valor é R$ {br(resumo['pib_per_capita_minimo'])} (Rafael Fernandes) e o maior "
        f"é R$ {br(resumo['pib_per_capita_maximo'])}. A média estadual é R$ {br(resumo['pib_per_capita_media_rn'])} "
        f"e a mediana é R$ {br(resumo['pib_per_capita_mediana_rn'])}. A distância entre média e mediana "
        "mostra assimetria à direita: poucos municípios com PIB per capita muito alto puxam a média para cima."
    )
    hist = _grafico_histograma(df, resumo["percentil_10_pib_per_capita"])
    pdf.image(hist, w=182)
    pdf.ln(2)
    pdf.corpo(
        "A tabela abaixo traz o ranking completo, do menor para o maior PIB per capita. "
        "A coluna “< P10” indica os municípios que entram no grupo da questão 2.2."
    )

    tabela_q1 = df.sort_values("ranking_pib_per_capita_asc")
    linhas_q1 = []
    for _, row in tabela_q1.iterrows():
        linhas_q1.append(
            [
                str(int(row["ranking_pib_per_capita_asc"])),
                row["municipio"],
                row["mesorregiao"],
                br(row["pib_per_capita"]),
                "Sim" if bool(row["abaixo_percentil_10"]) else "Não",
            ]
        )
    pdf.tabela(
        ["#", "Município", "Mesorregião", "PIB per capita (R$)", "< P10"],
        linhas_q1,
        [12, 58, 52, 38, 22],
    )

    pdf.add_page()
    pdf.titulo("2. Grupo com os 10% municípios de menor PIB per capita")
    pdf.corpo(
        f"Dez por cento de {resumo['n_municipios']} municípios correspondem a 16,7 observações. "
        "A forma estatística de recortar esse décimo inferior é o percentil 10. "
        f"Com a regra “PIB per capita menor que o P10”, o grupo ficou com "
        f"{resumo['n_municipios_abaixo_p10']} municípios, o que equivale a cerca de 10,2% do estado."
    )

    pdf.titulo("2.1 Cálculo do percentil 10")
    pdf.corpo(
        "O percentil 10 foi calculado sobre o PIB per capita dos 167 municípios, com interpolação "
        "linear (numpy.percentile, método padrão). Interpretação: aproximadamente 10% dos municípios "
        "potiguares têm PIB per capita inferior a esse valor, e cerca de 90% têm valor igual ou superior."
    )
    pdf.destaque(f"Percentil 10 (P10) = R$ {br(resumo['percentil_10_pib_per_capita'])}")
    pdf.corpo(
        f"Média do grupo abaixo do P10: R$ {br(resumo['pib_per_capita_media_grupo_p10'])}. "
        f"Média dos demais municípios: R$ {br(resumo['pib_per_capita_media_demais'])}."
    )

    pdf.titulo("2.2 Municípios com PIB per capita menor que o percentil 10")
    pdf.corpo(
        "A regra adotada é estrita: entram no grupo apenas os municípios com PIB per capita < P10. "
        "Nenhum município ficou exatamente igual ao percentil 10. Os 17 municípios, em ordem crescente, são:"
    )
    grupo = df[df["abaixo_percentil_10"]].sort_values("pib_per_capita")
    linhas_g = []
    for _, row in grupo.iterrows():
        linhas_g.append(
            [
                str(int(row["ranking_pib_per_capita_asc"])),
                row["municipio"],
                row["microrregiao"],
                br(row["pib_per_capita"]),
                br(row["populacao_censo_2022"], 0),
            ]
        )
    pdf.tabela(
        ["#", "Município", "Microrregião", "PIB per capita (R$)", "População"],
        linhas_g,
        [12, 50, 52, 40, 28],
    )

    pdf.titulo("2.3 Variáveis que podem explicar esses PIBs per capita")
    pdf.corpo(
        "O PIB per capita é a razão entre o valor da produção municipal e a população. "
        "Por isso, o nível baixo no décimo inferior pode estar associado a porte demográfico "
        "reduzido, pouca diversificação produtiva, forte peso da administração pública e "
        "indicadores sociais mais frágeis. As variáveis abaixo foram reunidas para esse recorte. "
        "O valor adicionado setorial é de 2021, último ano em que o IBGE divulgou a composição "
        "completa; população, área, densidade e alfabetização vêm do Censo 2022."
    )
    pdf.corpo(
        "• População e densidade: municípios menores e mais dispersos tendem a gerar menos "
        "atividade econômica por habitante.\n"
        "• Área territorial: indica o porte físico e, junto com a densidade, o grau de dispersão.\n"
        "• Participação da administração pública no VAB: economia dependente de salários e "
        "serviços públicos, com pouca base privada.\n"
        "• Participação da indústria e dos serviços: sinal de diversificação e de encadeamentos produtivos.\n"
        "• Agropecuária: peso do setor primário.\n"
        "• Taxa de alfabetização e escolarização: escolaridade associada à produtividade.\n"
        "• IDHM: síntese de renda, longevidade e educação.\n"
        "• Salário médio e população ocupada: formalização e tamanho do mercado de trabalho."
    )

    setores = _grafico_setores(resumo)
    pdf.image(setores, w=182)
    pdf.ln(2)

    pdf.corpo("Comparativo de médias: grupo abaixo do P10 versus os demais municípios do RN.")
    linhas_c = []
    for item in resumo["comparativo_grupo_vs_demais"]:
        casas = 3 if item["variavel"] == "IDHM" else 2
        linhas_c.append(
            [
                item["variavel"],
                br(item["media_grupo_p10"], casas),
                br(item["media_demais_municipios"], casas),
            ]
        )
    pdf.tabela(
        ["Variável", "Média do grupo P10", "Média dos demais"],
        linhas_c,
        [92, 45, 45],
    )

    pdf.corpo("Valores municipais das principais variáveis explicativas no grupo abaixo do P10:")
    linhas_v = []
    for _, row in grupo.iterrows():
        linhas_v.append(
            [
                row["municipio"],
                br(row["participacao_adm_publica_pct"], 1),
                br(row["participacao_industria_pct"], 1),
                br(row["taxa_alfabetizacao_15_mais_pct"], 1),
                br(row["idhm"], 3),
            ]
        )
    pdf.tabela(
        ["Município", "Adm. púb. (%)", "Indústria (%)", "Alfabetização (%)", "IDHM"],
        linhas_v,
        [52, 32, 32, 38, 28],
    )

    pdf.titulo("Descrição sumária da situação")
    pdf.corpo(resumo["descricao"])
    pdf.corpo(
        "Em síntese: o décimo inferior do PIB per capita no RN não é um recorte aleatório. "
        "Ele reúne municípios pequenos do Oeste e do Agreste Potiguar, com economia pouco "
        "industrializada, alta dependência do setor público e indicadores sociais abaixo da "
        "média estadual. Isso é coerente com um PIB por habitante persistentemente baixo."
    )

    pdf.titulo("Fontes")
    pdf.corpo(
        "IBGE. Produto Interno Bruto dos Municípios 2023 (pesquisa 38, indicador 47001).\n"
        "IBGE. Censo Demográfico 2022, tabelas 4714 (população, área e densidade) e 9543 (alfabetização).\n"
        "IBGE. Valor adicionado bruto por atividade econômica, 2021 (pesquisa 38).\n"
        "PNUD/IPEA/FJP. Índice de Desenvolvimento Humano Municipal, via IBGE Cidades@.\n"
        "Portal de referência: https://cidades.ibge.gov.br/brasil/rn/panorama"
    )

    destino.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(destino))
    return destino


def _estilo_cabecalho(ws, colunas: int) -> None:
    preenchimento = PatternFill("solid", fgColor="1F4E79")
    fonte = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    alinhamento = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda = Border(
        left=Side(style="thin", color="BDD7EE"),
        right=Side(style="thin", color="BDD7EE"),
        top=Side(style="thin", color="BDD7EE"),
        bottom=Side(style="thin", color="BDD7EE"),
    )
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, colunas + 1):
        cel = ws.cell(1, col)
        cel.fill = preenchimento
        cel.font = fonte
        cel.alignment = alinhamento
        cel.border = borda


def _ajustar_larguras(ws) -> None:
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        tamanho = 12
        for cel in col[:80]:
            if cel.value is None:
                continue
            tamanho = max(tamanho, min(len(str(cel.value)) + 2, 42))
        ws.column_dimensions[letra].width = tamanho


def _gravar_aba(ws, dados: pd.DataFrame) -> None:
    for linha in dataframe_to_rows(dados, index=False, header=True):
        ws.append(linha)
    _estilo_cabecalho(ws, dados.shape[1])
    zebra = PatternFill("solid", fgColor="F2F5FA")
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        if i % 2 == 0:
            for cel in row:
                cel.fill = zebra
        for cel in row:
            cel.alignment = Alignment(vertical="center")
            if isinstance(cel.value, float):
                cel.number_format = "#,##0.00"
    _ajustar_larguras(ws)


def gerar_xlsx(df: pd.DataFrame, resumo: dict, destino: Path) -> Path:
    completo = df.sort_values("ranking_pib_per_capita_asc").copy()
    completo["abaixo_percentil_10"] = completo["abaixo_percentil_10"].map(
        {True: "Sim", False: "Não", "True": "Sim", "False": "Não"}
    )
    cols = [c for c in ORDEM_XLSX if c in completo.columns]
    completo_pt = completo[cols].rename(columns=COLUNAS_PT)
    grupo_pt = completo_pt[completo_pt["Abaixo do percentil 10"] == "Sim"].copy()

    comparativo = pd.DataFrame(resumo["comparativo_grupo_vs_demais"]).rename(
        columns={
            "variavel": "Variável",
            "media_grupo_p10": "Média do grupo abaixo do P10",
            "media_demais_municipios": "Média dos demais municípios",
            "mediana_grupo_p10": "Mediana do grupo abaixo do P10",
            "mediana_demais_municipios": "Mediana dos demais municípios",
        }
    )
    percentil = pd.DataFrame(
        [
            ["Estado", "Rio Grande do Norte"],
            ["Total de municípios", resumo["n_municipios"]],
            ["Ano do PIB per capita", resumo["ano_pib"]],
            ["Separatriz", "Percentil 10"],
            ["Valor do percentil 10 (R$)", resumo["percentil_10_pib_per_capita"]],
            ["Regra do grupo", "PIB per capita < percentil 10"],
            ["Municípios abaixo do P10", resumo["n_municipios_abaixo_p10"]],
            ["PIB per capita mínimo (R$)", resumo["pib_per_capita_minimo"]],
            ["PIB per capita máximo (R$)", resumo["pib_per_capita_maximo"]],
            ["Média do RN (R$)", round(resumo["pib_per_capita_media_rn"], 2)],
            ["Mediana do RN (R$)", resumo["pib_per_capita_mediana_rn"]],
            ["Média do grupo P10 (R$)", round(resumo["pib_per_capita_media_grupo_p10"], 2)],
            ["Média dos demais (R$)", round(resumo["pib_per_capita_media_demais"], 2)],
            ["Fonte", "IBGE Cidades@, pesquisa 38, indicador 47001"],
        ],
        columns=["Item", "Valor"],
    )
    dicionario = pd.DataFrame(
        [{"Coluna": COLUNAS_PT[k], "Campo original": k} for k in cols]
    )

    wb = Workbook()
    abas = [
        ("1_PIB_per_capita_RN", completo_pt),
        ("2.1_Percentil_10", percentil),
        ("2.2_Grupo_abaixo_P10", grupo_pt),
        ("2.3_Comparativo", comparativo),
        ("Dicionario", dicionario),
    ]
    ws0 = wb.active
    ws0.title = abas[0][0]
    _gravar_aba(ws0, abas[0][1])
    for nome, dados in abas[1:]:
        ws = wb.create_sheet(nome)
        _gravar_aba(ws, dados)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def gerar_entregas(df: pd.DataFrame | None = None, resumo: dict | None = None) -> dict[str, Path]:
    if df is None:
        csv = PASTA_DADOS / "municipios_rn.csv"
        df = pd.read_csv(csv)
        df["abaixo_percentil_10"] = df["abaixo_percentil_10"].astype(str).isin(["True", "true", "1"])
    if resumo is None:
        json_resumo = PASTA_DADOS / "resumo_separatriz.json"
        if json_resumo.exists():
            resumo = json.loads(json_resumo.read_text(encoding="utf-8"))
        else:
            resumo = resumo_separatriz(df)

    pdf = PASTA_DADOS / "Relatorio_Separatrizes_PIB_RN.pdf"
    xlsx = PASTA_DADOS / "Separatrizes_PIB_per_capita_RN.xlsx"
    gerar_pdf(df, resumo, pdf)
    gerar_xlsx(df, resumo, xlsx)
    return {"pdf": pdf, "xlsx": xlsx}
